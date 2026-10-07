"""Build the editable SkillCraft DBTHON 2026 panel presentation.

Requires python-pptx. The ER diagram PNG is generated from the adjacent SVG
and is checked into the repository so this script needs no rendering service.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
OUTPUT = ROOT / "SkillCraft_DBTHON_2026_Panel.pptx"
DIAGRAM = ROOT / "skillcraft-er-diagram.png"

PAPER = RGBColor(248, 247, 242)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(32, 54, 46)
GREEN = RGBColor(66, 106, 84)
MINT = RGBColor(228, 236, 228)
WARM = RGBColor(238, 231, 222)
RUST = RGBColor(168, 105, 72)
MUTED = RGBColor(91, 111, 99)
LINE = RGBColor(205, 216, 205)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rect(slide, x, y, w, h, fill=WHITE, stroke=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if stroke is None else None
    if stroke is not None:
        shape.line.color.rgb = stroke
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.06
    return shape


def txt(slide, text, x, y, w, h, size=18, color=INK, bold=False,
        font="Aptos", align=None, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.015)
    tf.margin_top = tf.margin_bottom = Inches(0.015)
    tf.vertical_anchor = valign
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(0)
        if align is not None:
            p.alignment = align
    return box


def rule(slide, x1, y1, x2, y2, color=LINE, width=1):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color
    s.line.width = Pt(width)
    return s


def pill(slide, text, x, y, w, fill=MINT, color=GREEN, size=11):
    rect(slide, x, y, w, 0.34, fill, radius=True)
    txt(slide, text, x + 0.08, y + 0.045, w - 0.16, 0.25, size, color, True)


def card(slide, x, y, w, h, label, body, accent=GREEN, fill=WHITE,
         label_size=17, body_size=15):
    rect(slide, x, y, w, h, fill, LINE, True)
    rect(slide, x, y, 0.07, h, accent)
    txt(slide, label, x + 0.26, y + 0.19, w - 0.50, 0.40, label_size, INK, True)
    txt(slide, body, x + 0.26, y + 0.70, w - 0.52, h - 0.80, body_size, MUTED)


def base(title, subtitle, section, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, PAPER)
    txt(slide, "SKILLCRAFT  /  DBTHON 2026", 0.55, 0.28, 5.4, 0.26, 10, GREEN, True)
    txt(slide, section.upper(), 10.0, 0.28, 2.73, 0.26, 10, MUTED, True, align=PP_ALIGN.RIGHT)
    txt(slide, title, 0.55, 0.70, 12.1, 0.58, 30, INK, True)
    if subtitle:
        txt(slide, subtitle, 0.57, 1.35, 12.0, 0.45, 15, MUTED)
    rule(slide, 0.55, 7.11, 12.78, 7.11)
    txt(slide, "SkillCraft  •  database systems prototype", 0.56, 7.17, 7, 0.20, 9, MUTED)
    txt(slide, f"{n:02d}", 12.12, 7.16, 0.60, 0.20, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, message):
    slide.notes_slide.notes_text_frame.text = message


# 1 — cover
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.333, 7.5, PAPER)
rect(s, 8.62, 0, 4.713, 7.5, MINT)
txt(s, "DBTHON 2026  /  DATABASE SYSTEMS HACKATHON", 0.66, 0.60, 7.6, 0.28, 11, GREEN, True)
txt(s, "SkillCraft", 0.64, 1.45, 7.2, 0.90, 46, INK, True)
txt(s, "Local work, with a reliable record.", 0.67, 2.50, 7.4, 0.72, 27, INK, font="Georgia")
txt(s, "Employers post jobs. Artisans bid. The database keeps every decision, agreement and review connected.", 0.69, 3.58, 6.95, 1.0, 20, MUTED)
pill(s, "WORKING PROTOTYPE", 0.69, 5.25, 2.20)
txt(s, "Vedika Agarwal  ·  Bhavyaveer Kumar  ·  Anuj Deshpande", 0.69, 6.25, 7.2, 0.35, 13, INK)
txt(s, "Faculty guide: Siva Sankari  ·  2026–27", 0.69, 6.65, 6.5, 0.28, 11, MUTED)
for y, number, label in [(1.30, "01", "A job is posted"), (2.73, "02", "An artisan is selected"), (4.16, "03", "A contract records the outcome")]:
    rect(s, 9.05, y, 3.67, 1.06, WHITE, LINE, True)
    txt(s, number, 9.30, y + 0.21, 0.50, 0.42, 20, GREEN, True)
    txt(s, label, 9.93, y + 0.24, 2.45, 0.58, 17, INK, True)
notes(s, "Open with one sentence: SkillCraft is a local-work marketplace whose main strength is the database record behind each job. Introduce the team.")

# 2 — problem and domain
s = base("The problem is missing work history", "Small jobs often live in calls and chat messages; decisions are hard to trace later.", "problem + domain · 4 marks", 2)
card(s, 0.60, 2.07, 3.77, 3.10, "Today: scattered messages", "Who applied? Which price was accepted? Was payment acknowledged? The answers can be buried or disputed.", RUST, WARM, body_size=18)
card(s, 4.78, 2.07, 3.77, 3.10, "SkillCraft: linked records", "The original job, every bid, the chosen artisan, the contract and the later review stay connected.", GREEN, WHITE, body_size=18)
card(s, 8.97, 2.07, 3.77, 3.10, "Who it serves", "Local artisans seeking visible opportunities, and employers comparing offers for short projects.", GREEN, WHITE, body_size=18)
pill(s, "DOMAIN: LOCAL EMPLOYMENT", 0.62, 5.64, 2.73)
txt(s, "SDG 8 alignment: access to decent work and a portable work record. Economic impact is a goal, not a measured result yet.", 3.64, 5.55, 8.55, 0.73, 17, MUTED)
notes(s, "Explain the old approach first: informal messages do not create a dependable sequence. Then state the target users and SDG 8 connection. Do not claim proven earnings gains.")

# 3 — workflow
s = base("One job becomes one traceable agreement", "The same workflow is visible from two different accounts.", "working workflow", 3)
steps = [
    ("1  POST", "Employer adds a brief, location and budget."),
    ("2  DISCOVER", "Artisan searches by skill and place."),
    ("3  BID", "Artisan offers a price and optional note."),
    ("4  SELECT", "Employer accepts one bid; other bids close."),
    ("5  ACKNOWLEDGE", "Employer records payment sent; artisan confirms receipt."),
    ("6  REVIEW", "Paid contract receives a review and updates trust score."),
]
for i, (title, body) in enumerate(steps):
    x = 0.62 + (i % 3) * 4.23
    y = 2.02 + (i // 3) * 2.22
    card(s, x, y, 3.82, 1.83, title, body, GREEN if i not in (3, 4) else RUST, body_size=15)
notes(s, "Use Ramesh and Lakshmi as the concrete example. Emphasize that accepting a bid creates exactly one contract and that paid status for a new contract needs both sides.")

# 4 — role views
s = base("Each account sees the work relevant to it", "The roles share the same underlying job and contract, but their menus and actions differ.", "prototype", 4)
card(s, 0.70, 2.03, 5.76, 3.65, "EMPLOYER  /  RAMESH", "Hiring overview\nPost a job and inspect bids\nAccept or decline an offer\nRecord payment sent\nReview a paid contract", RUST, WARM, 20, 17)
card(s, 6.87, 2.03, 5.76, 3.65, "ARTISAN  /  LAKSHMI", "My workspace\nFind work and submit a bid\nTrack bid status\nConfirm payment received\nSee agreements and reviews", GREEN, MINT, 20, 17)
txt(s, "The API also checks the role and ownership; the database applies row-level rules for private bids, contracts and events.", 1.02, 6.15, 11.45, 0.47, 16, MUTED, align=PP_ALIGN.CENTER)
notes(s, "Show the actual role navigation in the live prototype. Ramesh cannot use artisan bid actions; Lakshmi cannot post jobs or inspect someone else's bids.")

# 5 — stack
s = base("The database does the critical work", "A plain-language map of the technology stack.", "DBMS implementation · 5 marks", 5)
for x, label, desc, bg in [
    (0.78, "React + Vite", "Website pages and forms", WHITE),
    (4.96, "Express API", "Login and request checks", WARM),
    (9.14, "PostgreSQL", "Records, rules and queries", MINT),
]:
    rect(s, x, 2.43, 3.44, 2.31, bg, LINE, True)
    txt(s, label, x + 0.25, 2.83, 2.98, 0.50, 22, INK, True, align=PP_ALIGN.CENTER)
    txt(s, desc, x + 0.24, 3.57, 2.99, 0.67, 17, MUTED, align=PP_ALIGN.CENTER)
txt(s, "→", 4.38, 3.12, 0.46, 0.56, 29, GREEN, True)
txt(s, "→", 8.55, 3.12, 0.46, 0.56, 29, GREEN, True)
rect(s, 1.30, 5.22, 10.75, 0.81, WHITE, LINE, True)
txt(s, "Local preview: PGlite runs the PostgreSQL-compatible schema. Docker configuration uses PostgreSQL 15.", 1.61, 5.46, 10.14, 0.34, 16, MUTED, align=PP_ALIGN.CENTER)
notes(s, "React is the visible website. Express is the server API. PostgreSQL is the relational database that joins records and enforces integrity. PGlite is the local preview database, not a separate product design.")

# 6 — ER diagram. Give the detailed model the whole slide so fields remain legible.
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.333, 7.5, PAPER)
s.shapes.add_picture(str(DIAGRAM), Inches(0.10), Inches(0.10), width=Inches(13.13), height=Inches(7.30))
notes(s, "Point to USERS, then GIG_POSTINGS, GIG_APPLICATIONS, COMPLETION_CONTRACTS, and finally RATINGS_REVIEWS plus CONTRACT_EVENTS. The accepted-bid trigger creates the contract; there is no direct bid-to-contract foreign key.")

# 7 — integrity
s = base("Rules protect the data when people act at once", "These protections are in SQL and transactions, not only in button logic.", "technical depth · 5 marks", 7)
cards = [
    (0.64, 2.01, "Keys + checks", "One bid per artisan per job; one contract per job; positive amounts; valid statuses."),
    (6.76, 2.01, "Safe acceptance", "A transaction locks the job row. Accepting one bid closes the job and rejects competitors."),
    (0.64, 4.12, "Triggers + functions", "Contract creation, event history and trust-score refresh happen in the database."),
    (6.76, 4.12, "Two confirmations", "A new contract cannot be marked paid until payment-sent and receipt timestamps both exist."),
]
for x, y, h, body in cards:
    card(s, x, y, 5.92, 1.74, h, body, GREEN, body_size=16)
notes(s, "One of the strongest demonstrations is two attempts to accept different bids on the same job. Exactly one succeeds. Do not attribute the result to a single mechanism; the lock, transaction, trigger and unique constraint work together.")

# 8 — innovation / novelty
s = base("The innovation is a verifiable sequence of claims", "Conventional status flags say 'paid'; SkillCraft records who said what and when.", "innovation 4 + novelty 5", 8)
rect(s, 0.70, 2.00, 5.48, 3.13, WARM, LINE, True)
txt(s, "CONVENTIONAL APPROACH", 1.00, 2.28, 4.78, 0.36, 15, RUST, True)
txt(s, "Employer flips one\npayment flag to paid.", 1.00, 2.88, 4.86, 1.10, 25, INK, True)
txt(s, "The worker's acknowledgment is missing.", 1.00, 4.34, 4.80, 0.47, 17, MUTED)
rect(s, 7.12, 2.00, 5.48, 3.13, MINT, LINE, True)
txt(s, "SKILLCRAFT NOW", 7.42, 2.28, 4.78, 0.36, 15, GREEN, True)
txt(s, "Employer records sent.\nArtisan confirms received.", 7.42, 2.88, 4.86, 1.10, 25, INK, True)
txt(s, "An append-only history records each event.", 7.42, 4.34, 4.80, 0.47, 17, MUTED)
txt(s, "A review is allowed only after a new contract is paid. Older paid rows remain visibly legacy, not falsely backfilled.", 1.05, 5.66, 11.25, 0.75, 17, MUTED, align=PP_ALIGN.CENTER)
notes(s, "This is the main novelty story. Acknowledgments are assertions by users, not bank verification. Explain why old paid rows are labelled legacy: the migration refuses to invent an artisan confirmation.")

# 9 — contract-linked trust
s = base("A review updates a careful trust score", "Only a paid contract can receive a review. A trigger refreshes the artisan's stored score.", "database function + trigger", 9)
rect(s, 0.75, 2.03, 11.83, 2.16, MINT, LINE, True)
txt(s, "ONE NEW 5-STAR REVIEW", 1.07, 2.33, 5.0, 0.35, 15, GREEN, True)
txt(s, "(5 + 3 × 3.5) ÷ (1 + 3)  =  3.88", 1.08, 2.92, 11.09, 0.71, 29, INK, True, align=PP_ALIGN.CENTER)
card(s, 0.78, 4.62, 3.72, 1.43, "ELIGIBLE DATA", "Reviews tied to paid contracts.", GREEN, WHITE, 17, 15)
card(s, 4.81, 4.62, 3.72, 1.43, "STABILIZING PRIOR", "Three virtual 3.5-star reviews.", RUST, WARM, 17, 15)
card(s, 8.84, 4.62, 3.72, 1.43, "AUTOMATIC UPDATE", "Review insert runs the score function.", GREEN, WHITE, 17, 15)
notes(s, "Explain the prior without jargon: one perfect review should not immediately make a new artisan look perfect. The stored score is recalculated by a database trigger after a real eligible review. This is a design choice, not a validated fairness guarantee.")

# 10 — search + security
s = base("Matching and privacy are database features", "The discovery screens are backed by indexed queries and restricted database access.", "implementation + differentiation", 10)
card(s, 0.73, 2.05, 3.77, 3.43, "SEARCH", "Jobs and artisans can be filtered by skill and searched by place or keyword. B-tree and full-text indexes support those queries.", GREEN, WHITE, 19, 17)
card(s, 4.78, 2.05, 3.77, 3.43, "EXPLAINABLE FIT", "A simple 0–3 score prefers the same skill (2 points) and same city (1 point). It is inspectable, not an AI prediction.", RUST, WARM, 19, 17)
card(s, 8.84, 2.05, 3.77, 3.43, "ROW SECURITY", "A restricted database role and row-level policies limit access to a person's private bids, contracts and events.", GREEN, MINT, 19, 17)
txt(s, "Public artisan reputation appears through limited views; private account fields stay private.", 1.03, 6.00, 11.2, 0.45, 17, MUTED, align=PP_ALIGN.CENTER)
notes(s, "Use the search screen to show a match label. Say the ranking is transparent, not AI. Explain row-level security in ordinary words: the database itself filters rows based on the signed-in user's identity.")

# 11 — validation
s = base("What we have actually validated", "Functional evidence is strong; performance and social impact still need measurement.", "validation · 3 marks", 11)
card(s, 0.72, 2.01, 3.76, 2.94, "API + SCHEMA", "Checks the full job-to-review flow, foreign-key relationships, audit immutability, and ownership rules.", GREEN, WHITE, 19, 17)
card(s, 4.78, 2.01, 3.76, 2.94, "24 JEST CASES", "20 CRUD cases plus 4 concurrency and integrity cases passed on fresh local PGlite databases.", GREEN, WHITE, 19, 17)
card(s, 8.84, 2.01, 3.76, 2.94, "2 BROWSER FLOWS", "The two-role payment journey and long-text mobile layout passed in Edge.", GREEN, WHITE, 19, 17)
rect(s, 1.22, 5.45, 10.89, 0.83, WARM, LINE, True)
txt(s, "Next evidence: benchmark indexed search vs. a baseline on native PostgreSQL; measure fit quality with labeled examples.", 1.53, 5.65, 10.27, 0.42, 16, INK, align=PP_ALIGN.CENTER)
notes(s, "Do not present test counts as a speed-up. These tests show behavior and integrity. Native PostgreSQL query-plan and ranking-quality comparisons are still planned.")

# 12 — rubric 1–4
s = base("What the panel grades · part 1", "The numbers are available marks, not a score awarded to this project.", "rubric / 30 marks", 12)
rubric_a = [
    ("4", "Problem + domain", "Informal local hiring lacks a reliable, portable record."),
    ("5", "Database design", "Seven linked tables, ER diagram, keys, constraints and normalization."),
    ("5", "DBMS depth", "Transactions, triggers, functions, indexes, views and row-level security."),
    ("4", "Innovation", "Two-party acknowledgment plus append-only contract events."),
]
for i, (marks, title, body) in enumerate(rubric_a):
    y = 1.94 + i * 1.16
    rect(s, 0.76, y, 11.81, 0.96, WHITE, LINE, True)
    rect(s, 1.00, y + 0.18, 0.67, 0.57, MINT, radius=True)
    txt(s, marks, 1.13, y + 0.28, 0.41, 0.31, 18, GREEN, True, align=PP_ALIGN.CENTER)
    txt(s, title, 1.98, y + 0.17, 3.01, 0.42, 19, INK, True)
    txt(s, body, 5.13, y + 0.17, 7.03, 0.63, 15, MUTED)
notes(s, "Tie every rubric criterion to one visible artifact. State that these are maximum marks allocated by the challenge PDF, not an estimated grade.")

# 13 — rubric 5–8
s = base("What the panel grades · part 2", "The claim and the evidence should stay at the same level.", "rubric / 30 marks", 13)
rubric_b = [
    ("5", "Novelty", "Compare with employer-only payment and a plain rating average."),
    ("2", "SDG impact", "SDG 8 goal: fairer access and a portable history; impact unmeasured."),
    ("3", "Validation", "Show passing tests; be honest that benchmark data is pending."),
    ("2", "TRL + demo", "Run the two-account prototype live; no field pilot yet."),
]
for i, (marks, title, body) in enumerate(rubric_b):
    y = 1.94 + i * 1.16
    rect(s, 0.76, y, 11.81, 0.96, WHITE, LINE, True)
    rect(s, 1.00, y + 0.18, 0.67, 0.57, MINT, radius=True)
    txt(s, marks, 1.13, y + 0.28, 0.41, 0.31, 18, GREEN, True, align=PP_ALIGN.CENTER)
    txt(s, title, 1.98, y + 0.17, 3.01, 0.42, 19, INK, True)
    txt(s, body, 5.13, y + 0.17, 7.03, 0.63, 15, MUTED)
notes(s, "The last two criteria are easy to overstate. Describe SDG alignment as an intention; describe TRL as a working prototype rather than a deployed or field-tested system.")

# 14 — demo guide
s = base("A five-minute live demonstration", "Use the connected local preview, not the browser-only demo mode.", "panel runbook", 14)
demo = [
    ("1", "Ramesh posts", "ramesh@constructionco.in"),
    ("2", "Lakshmi searches + bids", "lakshmi@skillcraft.local"),
    ("3", "Ramesh accepts", "One job closes; one contract appears"),
    ("4", "Both acknowledge", "Sent → received → paid"),
    ("5", "Ramesh reviews", "Show event history and trust score"),
]
for i, (number, heading, body) in enumerate(demo):
    y = 1.90 + i * 0.91
    rect(s, 0.75, y, 11.84, 0.75, WHITE, LINE, True)
    txt(s, number, 1.03, y + 0.16, 0.35, 0.36, 19, GREEN, True)
    txt(s, heading, 1.65, y + 0.14, 4.20, 0.39, 18, INK, True)
    txt(s, body, 6.00, y + 0.16, 5.97, 0.41, 15, MUTED)
pill(s, "BOTH PASSWORDS: password123", 0.76, 6.62, 3.14)
notes(s, "Prepare a fresh job before the demo and use its distinctive description to search. Avoid posting a new bid on a previously bid job. Both accounts use password123 in the local seeded preview.")

# 15 — limits / closing
s = base("A working prototype with clear next measurements", "The database is the product's backbone; our claims stop where the evidence stops.", "conclusion", 15)
card(s, 0.77, 2.08, 5.66, 3.20, "WHAT WORKS NOW", "Connected two-role website\nSeven-table relational model\nOne-winner bid acceptance\nTwo-party payment record\nContract-linked trust and audit history", GREEN, MINT, 19, 17)
card(s, 6.88, 2.08, 5.66, 3.20, "WHAT WE WILL MEASURE NEXT", "Native PostgreSQL query latency\nMatching quality on labeled jobs\nSecurity checks on a native server\nUser outcomes in a real pilot", RUST, WARM, 19, 17)
txt(s, "SkillCraft records payment claims. It does not send money, verify banks, or prove an earnings uplift.", 1.18, 5.81, 11.00, 0.58, 17, MUTED, align=PP_ALIGN.CENTER)
notes(s, "Close by restating the database contribution. Invite questions. If asked about payment, distinguish acknowledgment from real bank verification. If asked about performance, explain the planned benchmark rather than inventing a result.")

prs.save(OUTPUT)
print(f"Wrote {OUTPUT} ({len(prs.slides)} slides)")
